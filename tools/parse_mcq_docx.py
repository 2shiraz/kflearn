import docx,re,json,sys,collections
def parse(n):
    d=docx.Document(f'/mnt/user-data/uploads/MBBS{n}_MCQ_Bank_1000.docx')
    p=[x.text.strip() for x in d.paragraphs if x.text.strip()]
    # section headers: "MODULE k"/"BLOCK k" followed by name
    blocks=[];cur=None;topic=None;q=None
    body_end=next(k for k,x in enumerate(p) if k>20 and 'ANSWER KEY' in x.upper())
    covers=[re.sub(r'^\d+\. |\s+—\s+\d+ MCQs$','',x) for x in p[:12] if re.search(r'—\s+200 MCQs$',x)]
    k=0
    while k<body_end:
        x=p[k]
        if re.fullmatch(r'(MODULE|BLOCK) \d+',x) and k>8:
            cur={'name':covers[len(blocks)],'topics':[]};blocks.append(cur);topic=None;k+=2;continue
        m=re.fullmatch(r'(\d+)\. (.+?)(?:\s+—\s+\d+ MCQs)?',x)
        if cur and m and not x.endswith('?') and k+1<body_end and re.match(r'\d+\. ',p[k+1]) and k+2<body_end and p[k+2].startswith('A) '):
            topic={'name':m.group(2),'questions':[]};cur['topics'].append(topic);k+=1;continue
        if topic is not None:
            m=re.match(r'(\d+)\. (.+)',x)
            if m and k+1<body_end and p[k+1].startswith('A) '):
                q={'n':int(m.group(1)),'stem':m.group(2),'options':[]};topic['questions'].append(q);k+=1;continue
            m=re.match(r'([A-E])\) (.*)',x)
            if m and q: q['options'].append(m.group(2));k+=1;continue
        k+=1
    # topics listed in "Topics in this module" also match pattern—filter empties
    for b in blocks: b['topics']=[t for t in b['topics'] if t['questions']]
    topics=[t for b in blocks for t in b['topics']]
    key={}
    for tab in d.tables:
        for r in tab.rows[1:]:
            c=[x.text.strip() for x in r.cells]
            if c[0].isdigit(): key[int(c[0])]=(c[1],c[2])
    gi=0
    for t in topics:
        for q in t['questions']:
            gi+=1; q['n']=gi
    nums=[q['n'] for t in topics for q in t['questions']]
    assert nums==list(range(1,1001))
    assert sorted(key)==nums
    for t in topics:
        for q in t['questions']:
            a,r=key[q['n']]; q['answer']='ABCDE'.index(a); q['explanation']=r
    return blocks
out={}
for n in range(1,5):
    b=parse(n); out[n]=b
    c=collections.Counter(q['answer'] for x in b for t in x['topics'] for q in t['questions'])
    bad=[q['stem'][:40] for x in b for t in x['topics'] for q in t['questions'] if len(q['options'])!=5]
    print(n,[ (x['name'],len(x['topics']),sum(len(t['questions']) for t in x['topics'])) for x in b],dict(c),'bad',len(bad))
json.dump(out,open('all.json','w'))
